import sys
sys.stdout.reconfigure(encoding='utf-8')
import os
import re
import json
import sqlite3
import unicodedata
from pathlib import Path
import openpyxl
import pdfplumber
from rapidfuzz import fuzz

def strip_accents(s: str) -> str:
    if not s:
        return ""
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').upper()

# Medical canonical replacement rules
SYNONYMS = [
    # Citometria / Biometria
    (r'\bBIOMETRIA\s+HEMATICA\b', 'CITOMETRIA HEMATICA'),
    (r'\bB\.?H\.?\b', 'CITOMETRIA HEMATICA'),
    (r'\bCITOMETRIA\s+HEMATICA\s+COMPLETA\b', 'CITOMETRIA HEMATICA'),
    
    # Orina
    (r'\bEXAMEN\s+GENERAL\s+DE\s+ORINA\b', 'EXAMEN GENERAL DE ORINA'),
    (r'\bE\.?G\.?O\.?\b', 'EXAMEN GENERAL DE ORINA'),
    
    # Quimica sanguinea / Clinica / Integral
    (r'\bQUIMICA\s+CLINICA\b', 'QUIMICA SANGUINEA'),
    (r'\bQUIMICA\s+INTEGRAL\b', 'QUIMICA SANGUINEA'),
    (r'\bSUPER\s+QUIMICA\b', 'QUIMICA SANGUINEA'),
    (r'\bQUIMICA\s+DE\s+(\d+)\s+ELEMENTOS\b', r'QUIMICA SANGUINEA \1 ELEMENTOS'),
    (r'\bQUIMICA\s+(\d+)\b', r'QUIMICA SANGUINEA \1 ELEMENTOS'),
    (r'\bQS\s*(\d+)\b', r'QUIMICA SANGUINEA \1 ELEMENTOS'),
    (r'\bQS\b', 'QUIMICA SANGUINEA'),

    # Hemoglobina
    (r'\bHEMOGLOBINA\s+GLICOSILADA\b', 'HEMOGLOBINA GLICOSILADA HBA1C'),
    (r'\bHEMOGLOBINA\s+GLICADA\b', 'HEMOGLOBINA GLICOSILADA HBA1C'),
    (r'\bHBA1C\b', 'HEMOGLOBINA GLICOSILADA HBA1C'),

    # Prostatico
    (r'\bANTIGENO\s+PROSTATICO\s+ESPECIFICO\b', 'ANTIGENO PROSTATICO PSA'),
    (r'\bANTIGENO\s+PROSTATICO\b', 'ANTIGENO PROSTATICO PSA'),
    (r'\bPSA\s+TOTAL\b', 'ANTIGENO PROSTATICO PSA TOTAL'),
    (r'\bPSA\s+LIBRE\b', 'ANTIGENO PROSTATICO PSA LIBRE'),

    # Tiroides
    (r'\bPERFIL\s+TIROIDEO\s+COMPLETO\b', 'PERFIL TIROIDEO'),
    (r'\bPERFIL\s+TIROIDEO\s+I\b', 'PERFIL TIROIDEO'),
    
    # Lipidos
    (r'\bPERFIL\s+DE\s+LIPIDOS\b', 'PERFIL LIPIDICO'),
    (r'\bPERFIL\s+LIPIDOS\b', 'PERFIL LIPIDICO'),

    # Hepatico
    (r'\bPRUEBAS\s+DE\s+FUNCION\s+HEPATICA\b', 'PERFIL HEPATICO'),
    (r'\bPFH\b', 'PERFIL HEPATICO'),

    # Proteina C Reactiva
    (r'\bPROTEINA\s+C\s+REACTIVA\s+ULTRASENSIBLE\b', 'PROTEINA C REACTIVA ALTA SENSIBILIDAD'),
    (r'\bPROTEINA\s+C\s+REACTIVA\b', 'PROTEINA C REACTIVA'),
    (r'\bPCR\b', 'PROTEINA C REACTIVA'),

    # Vitamina D
    (r'\bVITAMINA\s+D\s*\(?25[\-\s]?OH\)?\s*(?:TOTAL)?\b', 'VITAMINA D 25 HIDROXI TOTAL'),
    (r'\b25\s+HIDROXI\s+VITAMINA\s+D\s+TOTAL\b', 'VITAMINA D 25 HIDROXI TOTAL'),

    # Imagenologia
    (r'\bRAYOS\s+X\b', 'RADIOGRAFIA RX'),
    (r'\bRX\b', 'RADIOGRAFIA RX'),
    (r'\bULTRASONIDO\b', 'ULTRASONIDO US'),
    (r'\bECOGRAFIA\b', 'ULTRASONIDO US'),
    (r'\bELECTROCARDIOGRAMA\s+EN\s+REPOSO\b', 'ELECTROCARDIOGRAMA ECG'),
    (r'\bECG\b', 'ELECTROCARDIOGRAMA ECG'),
    (r'\bELECTROCARDIOGRAMA\b', 'ELECTROCARDIOGRAMA ECG'),
]

# Common laboratory suffixes / noise
SPECIMEN_PATTERNS = [
    r'\bEN\s+SUERO\b',
    r'\bEN\s+SANGRE\b',
    r'\bEN\s+PLASMA\b',
    r'\bEN\s+ORINA\b',
    r'\bEN\s+LCR\b',
    r'\bEN\s+L\.C\.R\.\b',
    r'\bAL\s+AZAR\b',
    r'\bALEATORIA\b',
    r'\(Q\)',
    r'\(IFI\)',
    r'\(ELISA\)',
    r'\(CLIA\)',
    r'\(ECLIA\)',
]

def clean_medical_text(text: str, remove_specimen: bool = False) -> str:
    t = strip_accents(text)
    # clean extra characters
    t = re.sub(r'[\(\)\[\]\,\.\-\/\:\;\*\+\t\r\n]', ' ', t)
    t = ' '.join(t.split())
    # apply synonyms
    for pattern, repl in SYNONYMS:
        t = re.sub(pattern, repl, t)
    if remove_specimen:
        for p in SPECIMEN_PATTERNS:
            t = re.sub(p, ' ', t)
    t = ' '.join(t.split())
    return t

def extract_qs_elements(text: str):
    m = re.search(r'QUIMICA\s+SANGUINEA\s+(\d+)\s+ELEMENTOS', text)
    if m:
        return int(m.group(1))
    m2 = re.search(r'(\d+)\s*ELEMENTOS', text)
    if m2:
        return int(m2.group(1))
    return None

def build_matching_database():
    base_dir = Path(__file__).parent.parent
    db_path = base_dir / "db" / "chopo_prices.db"
    data_dir = base_dir / "data" / "lcm"
    pdf_path = data_dir / "ListaGeneral2025 actualizado a marzo de 2025.pdf"
    xlsx_path = data_dir / "2026 LCM PROMOCIONES octubre.xlsx"
    out_json = base_dir / "data" / "lcm" / "lcm_chopo_matches.json"

    # 1. Load Chopo catalog
    conn = sqlite3.connect(str(db_path))
    cur = conn.cursor()
    cur.execute("""
        SELECT study_name, price, price_original, url, sku
        FROM latest_prices
    """)
    chopo_raw = cur.fetchall()
    conn.close()

    chopo_catalog = []
    for row in chopo_raw:
        s_name = row[0]
        norm = clean_medical_text(s_name, remove_specimen=False)
        core = clean_medical_text(s_name, remove_specimen=True)
        elem = extract_qs_elements(norm)
        chopo_catalog.append({
            'name': s_name,
            'norm': norm,
            'core': core,
            'elements': elem,
            'price_web': row[1],
            'price_list': row[2],
            'url': row[3],
            'sku': row[4]
        })

    print(f"Loaded {len(chopo_catalog)} Chopo studies from SQLite.")

    # 2. Parse LCM PDF
    lcm_studies = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for p_idx, page in enumerate(pdf.pages):
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    if not row or len(row) < 3:
                        continue
                    clave, estudio, precio = row[0], row[1], row[2]
                    if clave == 'CLAVE' or not estudio or not precio:
                        continue
                    p_str = re.sub(r'[\s\$]', '', str(precio)).replace(',', '')
                    try:
                        p_val = float(p_str)
                        name_str = ' '.join(str(estudio).split())
                        norm = clean_medical_text(name_str, remove_specimen=False)
                        core = clean_medical_text(name_str, remove_specimen=True)
                        elem = extract_qs_elements(norm)
                        lcm_studies.append({
                            'code': str(clave).strip(),
                            'name': name_str,
                            'norm': norm,
                            'core': core,
                            'elements': elem,
                            'price': p_val,
                            'page': p_idx + 1
                        })
                    except ValueError:
                        pass

    print(f"Parsed {len(lcm_studies)} LCM studies from PDF.")

    # 3. Parse LCM Excel Packages & Promotions
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    lcm_packages = []
    ws1 = wb['OCTUBRE 2026']
    
    current_cat = "Promociones"
    for r in ws1.iter_rows(values_only=True):
        if r[0] and str(r[0]).strip():
            current_cat = str(r[0]).strip().replace('\n', ' ')
        p_name = r[1]
        if not p_name or not str(p_name).strip():
            continue
        p_name = str(p_name).strip()
        if p_name in ('PAQUETE', 'CUMPLEAÑEROS DEL MES DE OCTUBRE'):
            continue
            
        code = r[2] if r[2] else ""
        included = r[3] if r[3] else ""
        price_pub = r[4]
        price_disc = r[5]
        price_promo = r[6]
        
        # Parse promo price numeric
        promo_val = None
        if price_promo:
            m = re.search(r'\$?([0-9]+(?:,[0-9]{3})*(?:\.[0-9]+)?)', str(price_promo).replace(' ', ''))
            if m:
                promo_val = float(m.group(1).replace(',', ''))
        if promo_val is None and isinstance(price_promo, (int, float)):
            promo_val = float(price_promo)

        lcm_packages.append({
            'category': current_cat,
            'package_name': p_name,
            'code': str(code).strip(),
            'included_tests': str(included).strip(),
            'price_public': float(price_pub) if isinstance(price_pub, (int, float)) else None,
            'price_discount_system': float(price_disc) if isinstance(price_disc, (int, float)) else None,
            'price_promo': promo_val,
            'raw_promo_text': str(price_promo).strip() if price_promo else None
        })

    ws2 = wb['Adicionales 2026']
    lcm_adicionales = []
    for r in list(ws2.iter_rows(values_only=True))[3:]:
        if r[1] and r[2] is not None:
            lcm_adicionales.append({
                'code': str(r[0]).strip(),
                'name': str(r[1]).strip(),
                'price_2025': float(r[2]) if isinstance(r[2], (int, float)) else None,
                'price_bundle': float(r[3]) if isinstance(r[3], (int, float)) else None
            })

    print(f"Parsed {len(lcm_packages)} LCM packages and {len(lcm_adicionales)} adicionales from Excel.")

    # 4. Multi-stage intelligent matching
    # Stage A: Exact core match
    # Stage B: Element-matched QS
    # Stage C: Weighted fuzzy match on core + analyte
    matches = []
    
    for lcm in lcm_studies:
        lcm_norm = lcm['norm']
        lcm_core = lcm['core']
        lcm_elem = lcm['elements']
        
        # Exact Core match
        exact_c = None
        for c in chopo_catalog:
            if c['core'] == lcm_core:
                # If QS, check elements match
                if lcm_elem is not None and c['elements'] is not None:
                    if lcm_elem != c['elements']:
                        continue
                exact_c = c
                break
                
        if exact_c:
            price_chopo_web = exact_c['price_web']
            price_chopo_list = exact_c['price_list']
            diff_web = round(lcm['price'] - price_chopo_web, 2) if price_chopo_web else None
            diff_pct = round((diff_web / price_chopo_web) * 100, 1) if (diff_web is not None and price_chopo_web) else None
            matches.append({
                'lcm_code': lcm['code'],
                'lcm_name': lcm['name'],
                'lcm_price': lcm['price'],
                'chopo_name': exact_c['name'],
                'chopo_price_web': price_chopo_web,
                'chopo_price_list': price_chopo_list,
                'chopo_url': exact_c['url'],
                'chopo_sku': exact_c['sku'],
                'match_type': 'EXACT',
                'confidence': 100.0,
                'diff_mxn': diff_web,
                'diff_pct': diff_pct
            })
            continue

        # Filter candidates if QS
        candidates = chopo_catalog
        if lcm_elem is not None:
            filtered = [c for c in chopo_catalog if c['elements'] == lcm_elem]
            if filtered:
                candidates = filtered

        # Fuzzy matching
        best_score = 0
        best_cand = None
        
        for c in candidates:
            # Check element mismatch
            if lcm_elem is not None and c['elements'] is not None and lcm_elem != c['elements']:
                continue
                
            s1 = fuzz.token_sort_ratio(lcm_core, c['core'])
            s2 = fuzz.token_set_ratio(lcm_core, c['core'])
            # Penalize if one string is much longer and includes extra words
            len_ratio = min(len(lcm_core), len(c['core'])) / max(len(lcm_core), len(c['core'])) if max(len(lcm_core), len(c['core'])) > 0 else 1.0
            
            score = (s1 * 0.55) + (s2 * 0.35) + (len_ratio * 10)
            
            if score > best_score:
                best_score = score
                best_cand = c

        if best_cand and best_score >= 70:
            match_tier = 'HIGH' if best_score >= 85 else 'MEDIUM'
            p_web = best_cand['price_web']
            p_list = best_cand['price_list']
            diff_web = round(lcm['price'] - p_web, 2) if p_web else None
            diff_pct = round((diff_web / p_web) * 100, 1) if (diff_web is not None and p_web) else None
            
            matches.append({
                'lcm_code': lcm['code'],
                'lcm_name': lcm['name'],
                'lcm_price': lcm['price'],
                'chopo_name': best_cand['name'],
                'chopo_price_web': p_web,
                'chopo_price_list': p_list,
                'chopo_url': best_cand['url'],
                'chopo_sku': best_cand['sku'],
                'match_type': match_tier,
                'confidence': round(best_score, 1),
                'diff_mxn': diff_web,
                'diff_pct': diff_pct
            })
        else:
            matches.append({
                'lcm_code': lcm['code'],
                'lcm_name': lcm['name'],
                'lcm_price': lcm['price'],
                'chopo_name': None,
                'chopo_price_web': None,
                'chopo_price_list': None,
                'chopo_url': None,
                'chopo_sku': None,
                'match_type': 'NO_MATCH',
                'confidence': 0.0,
                'diff_mxn': None,
                'diff_pct': None
            })

    # Summary statistics
    exact_count = sum(1 for m in matches if m['match_type'] == 'EXACT')
    high_count = sum(1 for m in matches if m['match_type'] == 'HIGH')
    med_count = sum(1 for m in matches if m['match_type'] == 'MEDIUM')
    no_count = sum(1 for m in matches if m['match_type'] == 'NO_MATCH')
    total = len(matches)

    # Price competitiveness analysis on matched studies
    valid_comparisons = [m for m in matches if m['diff_mxn'] is not None]
    lcm_cheaper = sum(1 for m in valid_comparisons if m['diff_mxn'] < 0)
    chopo_cheaper = sum(1 for m in valid_comparisons if m['diff_mxn'] > 0)
    equal_price = sum(1 for m in valid_comparisons if m['diff_mxn'] == 0)

    output_payload = {
        'metadata': {
            'generated_at': '2026-09-30',
            'lcm_source_pdf': 'ListaGeneral2025 actualizado a marzo de 2025.pdf',
            'lcm_source_excel': '2026 LCM PROMOCIONES octubre.xlsx',
            'chopo_source_db': 'db/chopo_prices.db (Merida Altabrisa)',
            'total_lcm_studies': total,
            'exact_matches': exact_count,
            'high_confidence': high_count,
            'medium_confidence': med_count,
            'no_match': no_count,
            'actionable_matched_count': exact_count + high_count + med_count,
            'actionable_matched_pct': round((exact_count + high_count + med_count) / total * 100, 1),
            'price_analysis': {
                'total_compared': len(valid_comparisons),
                'lcm_cheaper_count': lcm_cheaper,
                'lcm_cheaper_pct': round(lcm_cheaper / len(valid_comparisons) * 100, 1) if valid_comparisons else 0,
                'chopo_cheaper_count': chopo_cheaper,
                'chopo_cheaper_pct': round(chopo_cheaper / len(valid_comparisons) * 100, 1) if valid_comparisons else 0,
                'equal_price_count': equal_price,
            }
        },
        'matches': matches,
        'packages': lcm_packages,
        'adicionales': lcm_adicionales
    }

    with open(str(out_json), 'w', encoding='utf-8') as f:
        json.dump(output_payload, f, ensure_ascii=False, indent=2)

    print("\n" + "="*60)
    print("RESUMEN DE MATCHING CLINICO LCM VS CHOPO MERIDA")
    print("="*60)
    print(f"Estudios totales en lista LCM: {total}")
    print(f"- Matches Exactos (100%): {exact_count} ({exact_count/total*100:.1f}%)")
    print(f"- Matches Alta Confianza (>=85%): {high_count} ({high_count/total*100:.1f}%)")
    print(f"- Matches Media Confianza (70-84%): {med_count} ({med_count/total*100:.1f}%)")
    print(f"- Sin match / catalogo exclusivo LCM: {no_count} ({no_count/total*100:.1f}%)")
    print(f"-> TOTAL CON COBERTURA DE COMPARACION: {exact_count + high_count + med_count} ({(exact_count + high_count + med_count)/total*100:.1f}%)")
    print("-" * 60)
    print(f"ANALISIS DE PRECIOS COMPARATIVOS ({len(valid_comparisons)} estudios):")
    print(f"- LCM es MAS BARATO en: {lcm_cheaper} estudios ({lcm_cheaper/len(valid_comparisons)*100:.1f}%)")
    print(f"- Chopo es MAS BARATO en: {chopo_cheaper} estudios ({chopo_cheaper/len(valid_comparisons)*100:.1f}%)")
    print(f"- Mismo precio exacto: {equal_price}")
    print("="*60)

if __name__ == '__main__':
    build_matching_database()
