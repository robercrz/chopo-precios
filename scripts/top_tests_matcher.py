import sys
sys.stdout.reconfigure(encoding='utf-8')
import sqlite3
import json
import re
import openpyxl
import pdfplumber
from rapidfuzz import fuzz

# Load Chopo
conn = sqlite3.connect('db/chopo_prices.db')
cur = conn.cursor()
cur.execute("SELECT study_name, price, price_original FROM latest_prices")
chopo_data = cur.fetchall()
conn.close()

# Load LCM PDF
pdf_file = r'data\lcm\ListaGeneral2025 actualizado a marzo de 2025.pdf'
lcm_pdf = {}
with pdfplumber.open(pdf_file) as pdf:
    for page in pdf.pages:
        tables = page.extract_tables()
        for t in tables:
            for r in t:
                if not r or len(r) < 3 or r[0] == 'CLAVE':
                    continue
                clave, name, pr = r[0], r[1], r[2]
                if not name or not pr:
                    continue
                clean_pr = re.sub(r'[\s\$]', '', str(pr)).replace(',', '')
                try:
                    val = float(clean_pr)
                    lcm_pdf[name.strip().upper()] = {'code': str(clave).strip(), 'name': name.strip(), 'price': val}
                except:
                    pass

# Load LCM Excel
wb = openpyxl.load_workbook(r'data\lcm\2026 LCM PROMOCIONES octubre.xlsx', data_only=True)
lcm_packages = []
ws1 = wb['OCTUBRE 2026']
for r in ws1.iter_rows(values_only=True):
    if r[1] and ('Check' in str(r[1]) or 'Perfil' in str(r[1]) or 'Rayos' in str(r[1]) or 'Ultrasonido' in str(r[1]) or 'ANT' in str(r[1]) or 'Vitamina' in str(r[1])):
        lcm_packages.append({'name': str(r[1]).strip(), 'price_promo': r[6], 'price_public': r[4]})

ws2 = wb['Adicionales 2026']
lcm_adicionales = []
for r in list(ws2.iter_rows(values_only=True))[3:]:
    if r[1] and r[2]:
        lcm_adicionales.append({'name': str(r[1]).strip(), 'code': str(r[0]), 'price_2025': r[2], 'price_bundle': r[3]})

print(f"Total LCM PDF studies: {len(lcm_pdf)}")
print(f"Total LCM Packages in Excel: {len(lcm_packages)}")
print(f"Total LCM Adicionales in Excel: {len(lcm_adicionales)}")

# Let's inspect Top 25 key tests
top_targets = [
    "BIOMETRIA HEMATICA",
    "EXAMEN GENERAL DE ORINA",
    "QUIMICA SANGUINEA",
    "GLUCOSA",
    "PERFIL DE LIPIDOS",
    "PERFIL TIROIDEO",
    "HEMOGLOBINA GLICOSILADA",
    "ANTIGENO PROSTATICO",
    "ACIDO URICO",
    "CREATININA",
    "UREA",
    "PERFIL HORMONAL",
    "COPROPARASITOSCOPICO",
    "TIEMPO DE PROTROMBINA",
    "GRUPO SANGUINEO",
    "REACCIONES FEBRILES",
    "VDRL",
    "VIH",
    "ELECTROLITOS",
    "VITAMINA D",
    "INSULINA",
    "FERRITINA",
    "UROCULTIVO",
    "PROTEINA C REACTIVA",
    "FACTOR REUMATOIDE"
]

print("\n=== TOP 25 TARGET SEARCH ===")
for target in top_targets:
    lcm_matches = [v for k, v in lcm_pdf.items() if target in k]
    if not lcm_matches:
        # try fuzzy
        for k, v in lcm_pdf.items():
            if fuzz.partial_ratio(target, k) > 85:
                lcm_matches.append(v)
    
    chopo_matches = [c for c in chopo_data if target in c[0].upper()]
    if not chopo_matches:
        for c in chopo_data:
            if fuzz.partial_ratio(target, c[0].upper()) > 85:
                chopo_matches.append(c)

    print(f"\n🎯 TARGET: {target}")
    print(f"   LCM ({len(lcm_matches)}):", [(m['name'], m['price']) for m in lcm_matches[:3]])
    print(f"   CHOPO ({len(chopo_matches)}):", [(c[0], c[1], c[2]) for c in chopo_matches[:3]])
